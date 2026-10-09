import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


EXPECTED_HEADERS = [
    "Record",
    "Animal",
    "Slice",
    "Thickness (mm)",
    "Contralateral area (mm²)",
    "Ipsilateral area (mm²)",
    "Raw infarct area (mm²)",
    "Corrected infarct area (mm²)",
    "Corrected infarct (%)",
    "Corrected volume (mm³)",
]


def read_rows(csv_path):
    with csv_path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream)
        rows = list(reader)

    if not rows or rows[0] != EXPECTED_HEADERS:
        raise ValueError("CSV columns do not match the TTC result format")

    converted = []
    for row in rows[1:]:
        if len(row) != len(EXPECTED_HEADERS):
            raise ValueError("CSV row does not contain 10 columns")
        converted.append(
            [
                int(row[0]),
                str(row[1]),
                row[2] if row[2] == "TOTAL" else int(row[2]),
                *[float(value) for value in row[3:]],
            ]
        )
    return converted


def build_workbook(rows, output_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TTC Results"
    sheet.sheet_view.showGridLines = False
    sheet.append(EXPECTED_HEADERS)
    for row in rows:
        sheet.append(row)

    dark_blue = "1F4E78"
    light_blue = "DCE6F1"
    body_text = "1F2937"
    white = "FFFFFF"

    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = Font(name="Arial", size=10, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 36

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=10):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color=body_text)
            cell.alignment = Alignment(vertical="center")
        row[0].alignment = Alignment(horizontal="center", vertical="center")
        row[1].alignment = Alignment(horizontal="center", vertical="center")
        row[2].alignment = Alignment(horizontal="center", vertical="center")
        for cell in row[3:]:
            cell.alignment = Alignment(horizontal="right", vertical="center")
        sheet.row_dimensions[row[0].row].height = 22

    for cell in sheet["A"][1:]:
        cell.number_format = "0"
    for cell in sheet["D"][1:]:
        cell.number_format = "0.0"
    for column in ("E", "F", "G", "H", "J"):
        for cell in sheet[column][1:]:
            cell.number_format = "0.000"
    for cell in sheet["I"][1:]:
        cell.number_format = '0.000"%"'

    total_row = sheet.max_row
    top_border = Side(style="medium", color=dark_blue)
    for cell in sheet[total_row]:
        cell.fill = PatternFill("solid", fgColor=light_blue)
        cell.font = Font(name="Arial", size=10, bold=True, color=body_text)
        cell.border = Border(top=top_border)

    widths = [9, 22, 10, 14, 22, 20, 21, 24, 20, 23]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    table = Table(displayName="TTCResults", ref=f"A1:J{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.auto_filter.ref = f"A1:J{sheet.max_row}"
    workbook.save(output_path)


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Usage: format_ttc_results.py input.csv output.xlsx")
    csv_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    rows = read_rows(csv_path)
    build_workbook(rows, output_path)
    print(f"OK: {output_path}")


if __name__ == "__main__":
    main()
